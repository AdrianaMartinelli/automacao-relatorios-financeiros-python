from pathlib import Path
import pandas as pd

BASE = Path(__file__).resolve().parents[1]
entrada = BASE / 'dados' / 'lancamentos.csv'
saida = BASE / 'saida' / 'relatorio_financeiro.xlsx'
saida.parent.mkdir(exist_ok=True)

df = pd.read_csv(entrada, sep=';', encoding='utf-8-sig')
obrigatorias = {'data', 'descricao', 'categoria', 'tipo', 'valor'}
if not obrigatorias.issubset(df.columns):
    raise ValueError(f'Colunas ausentes: {obrigatorias - set(df.columns)}')
df = df.drop_duplicates().copy()
df['data'] = pd.to_datetime(df['data'], format='%d/%m/%Y', errors='coerce')
df['valor'] = pd.to_numeric(df['valor'], errors='coerce')
df = df.dropna(subset=['data', 'valor', 'tipo'])
df = df[df['tipo'].isin(['Receita', 'Despesa']) & (df['valor'] >= 0)].copy()
df['mes'] = df['data'].dt.strftime('%Y-%m')
mensal = df.pivot_table(index='mes', columns='tipo', values='valor', aggfunc='sum', fill_value=0)
for coluna in ['Receita', 'Despesa']:
    if coluna not in mensal: mensal[coluna] = 0
mensal['Saldo'] = mensal['Receita'] - mensal['Despesa']
mensal = mensal.reset_index()
categorias = df.groupby(['tipo', 'categoria'], as_index=False)['valor'].sum()
with pd.ExcelWriter(saida, engine='openpyxl') as writer:
    df.to_excel(writer, sheet_name='Lancamentos', index=False)
    mensal.to_excel(writer, sheet_name='Resumo Mensal', index=False)
    categorias.to_excel(writer, sheet_name='Categorias', index=False)
    from openpyxl.chart import BarChart, Reference
    ws = writer.sheets['Resumo Mensal']
    chart = BarChart(); chart.title = 'Receitas e despesas por mês'; chart.y_axis.title = 'Valor (R$)'; chart.x_axis.title = 'Mês'
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(chart, 'F2')
    for sheet in writer.book:
        sheet.freeze_panes = 'A2'
        for cells in sheet.columns:
            letter = cells[0].column_letter
            sheet.column_dimensions[letter].width = min(35, max(13, max(len(str(c.value or '')) for c in cells) + 2))
print(f'Relatório gerado: {saida}')
